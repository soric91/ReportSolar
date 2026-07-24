import base64
import io
from datetime import date, datetime

import pytest
from docx.shared import Inches
from PIL import Image

from app.services import docx_generator as dg
from app.services.docx_generator import DocxGenerator, formatear_fecha


def _png_bytes():
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (10, 20, 30)).save(buf, "PNG")
    return buf.getvalue()


def _webp_bytes(size=(40, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, (10, 20, 30)).save(buf, "WEBP")
    return buf.getvalue()


class _Respuesta:
    def __init__(self, status_code, content=b""):
        self.status_code = status_code
        self.content = content


class _Settings:
    SUPABASE_URL = "https://x.supabase.co"
    SUPABASE_SERVICE_KEY = "service-key"


@pytest.fixture
def supabase_configurado(monkeypatch):
    monkeypatch.setattr(dg, "get_settings", lambda: _Settings())


class TestDescargaDeFotos:
    """La URL guardada es la pública del bucket y solo responde si el bucket
    está marcado público; por eso se descarga por path con la service key."""

    def test_prefiere_el_path_autenticado(self, monkeypatch, supabase_configurado):
        llamadas = []

        def fake_get(url, headers=None, **kwargs):
            llamadas.append((url, headers))
            return _Respuesta(200, _png_bytes())

        monkeypatch.setattr(dg.httpx, "get", fake_get)
        generator = DocxGenerator()
        datos = generator._descargar_imagen(
            {"url": "https://x.supabase.co/storage/v1/object/public/app_report/r/f.webp",
             "path": "r/f.webp"}
        )

        assert datos == _png_bytes()
        url, headers = llamadas[0]
        assert "/object/app_report/r/f.webp" in url and "/public/" not in url
        assert headers and "Authorization" in headers

    def test_cae_a_la_url_publica_si_el_path_falla(self, monkeypatch, supabase_configurado):
        llamadas = []

        def fake_get(url, headers=None, **kwargs):
            llamadas.append(url)
            if "/public/" not in url:
                return _Respuesta(400)
            return _Respuesta(200, _png_bytes())

        monkeypatch.setattr(dg.httpx, "get", fake_get)
        datos = DocxGenerator()._descargar_imagen(
            {"url": "https://x.supabase.co/storage/v1/object/public/app_report/r/f.webp",
             "path": "r/f.webp"}
        )

        assert datos == _png_bytes()
        assert len(llamadas) == 2

    def test_sin_path_usa_la_url(self, monkeypatch):
        monkeypatch.setattr(dg.httpx, "get", lambda url, **kw: _Respuesta(200, _png_bytes()))
        assert DocxGenerator()._descargar_imagen({"url": "https://x/f.webp"}) == _png_bytes()

    def test_data_uri_no_sale_a_la_red(self, monkeypatch):
        def explota(*a, **kw):
            raise AssertionError("no debería descargar un data URI")

        monkeypatch.setattr(dg.httpx, "get", explota)
        uri = "data:image/png;base64," + base64.b64encode(_png_bytes()).decode()
        assert DocxGenerator()._descargar_imagen({"url": uri}) == _png_bytes()

    def test_no_repite_la_descarga_de_la_misma_foto(self, monkeypatch):
        llamadas = []
        monkeypatch.setattr(
            dg.httpx, "get",
            lambda url, **kw: (llamadas.append(url), _Respuesta(200, _png_bytes()))[1],
        )
        generator = DocxGenerator()
        foto = {"url": "https://x/f.webp"}
        generator._descargar_imagen(foto)
        generator._descargar_imagen(foto)
        assert len(llamadas) == 1

    def test_devuelve_none_si_no_hay_nada(self):
        assert DocxGenerator()._descargar_imagen({}) is None


class TestNormalizarImagen:
    """El frontend sube las fotos en WebP y Word no lo admite:
    add_picture lanzaba UnrecognizedImageError y no entraba ninguna."""

    def test_convierte_webp_a_un_formato_que_word_admite(self):
        datos, _ = DocxGenerator._normalizar_imagen(_webp_bytes())
        assert Image.open(io.BytesIO(datos)).format in dg.FORMATOS_SOPORTADOS

    def test_deja_el_png_intacto(self):
        original = _png_bytes()
        datos, _ = DocxGenerator._normalizar_imagen(original)
        assert datos is original

    def test_informa_la_proporcion_para_poder_escalar(self):
        _, proporcion = DocxGenerator._normalizar_imagen(_webp_bytes((800, 400)))
        assert proporcion == pytest.approx(2.0)

    def test_una_foto_webp_termina_embebida(self, monkeypatch):
        monkeypatch.setattr(dg.httpx, "get", lambda url, **kw: _Respuesta(200, _webp_bytes()))
        generator = DocxGenerator()
        generator.add_section(
            {"titulo": "Módulos", "icono": "", "campos": []},
            {},
            [{"tipo": "antes", "url": "https://x/f.webp", "checklist_item": "modulos.Limpieza"}],
        )
        assert generator.doc.inline_shapes

    def test_la_foto_vertical_se_limita_por_alto(self, monkeypatch):
        """Escalada solo por ancho, una foto 709x1600 mide 4 pulgadas de alto
        y empuja la fila a la página siguiente."""
        monkeypatch.setattr(
            dg.httpx, "get", lambda url, **kw: _Respuesta(200, _webp_bytes((709, 1600)))
        )
        generator = DocxGenerator()
        generator.add_section(
            {"titulo": "Módulos", "icono": "", "campos": []},
            {},
            [{"tipo": "antes", "url": "https://x/f.webp", "checklist_item": "modulos.Limpieza"}],
        )
        forma = generator.doc.inline_shapes[0]
        assert forma.height <= Inches(1.7)
        assert forma.width < forma.height  # conserva la proporción vertical


class TestFotosPorCampo:
    """Agrupadas al final de la sección no se sabe a qué campo pertenecen"""

    SECCION = {
        "titulo": "Módulos",
        "icono": "",
        "campos": [{"nombre": "Limpieza de módulos"}, {"nombre": "Notas"}],
    }
    DATOS = {"Limpieza de módulos": "Se lavó", "Notas": "Un panel partido"}

    @staticmethod
    def _foto(campo, tipo):
        return {"tipo": tipo, "url": f"https://x/{campo}-{tipo}.webp",
                "checklist_item": f"modulos.{campo}"}

    def test_cada_campo_recibe_solo_sus_fotos(self):
        fotos = [
            self._foto("Limpieza de módulos", "antes"),
            self._foto("Limpieza de módulos", "despues"),
            self._foto("Notas", "foto_unica"),
        ]
        propias = DocxGenerator._fotos_del_campo(fotos, "Limpieza de módulos")
        assert [f["tipo"] for f in propias] == ["antes", "despues"]
        assert DocxGenerator._fotos_del_campo(fotos, "Notas") == [fotos[2]]

    def test_los_grupos_de_medidas_incluyen_las_de_cada_item(self):
        fotos = [{"tipo": "unica", "url": "https://x/a.webp",
                  "checklist_item": "medidas_dc.Voltajes por String_inv1_string1"}]
        assert DocxGenerator._fotos_del_campo(fotos, "Voltajes por String") == fotos

    def test_un_campo_no_se_queda_con_las_fotos_de_otro(self):
        fotos = [self._foto("Notas", "foto_unica")]
        assert DocxGenerator._fotos_del_campo(fotos, "Limpieza de módulos") == []

    def test_las_fotos_van_debajo_de_su_campo(self, monkeypatch):
        monkeypatch.setattr(dg.httpx, "get", lambda url, **kw: _Respuesta(200, _webp_bytes()))
        generator = DocxGenerator()
        generator.add_section(
            self.SECCION,
            self.DATOS,
            [self._foto("Limpieza de módulos", "antes"), self._foto("Notas", "foto_unica")],
        )

        # El orden de las tablas refleja el del documento: campo, sus fotos,
        # campo siguiente, sus fotos
        textos = [
            " ".join(c.text for f in t.rows for c in f.cells) for t in generator.doc.tables
        ]
        i_limpieza = next(i for i, t in enumerate(textos) if "Limpieza de módulos" in t)
        i_notas = next(i for i, t in enumerate(textos) if "Notas" in t)
        i_antes = next(i for i, t in enumerate(textos) if "Antes" in t)
        assert i_limpieza < i_antes < i_notas

    def test_una_foto_sin_campo_conocido_no_se_pierde(self, monkeypatch):
        monkeypatch.setattr(dg.httpx, "get", lambda url, **kw: _Respuesta(200, _webp_bytes()))
        generator = DocxGenerator()
        generator.add_section(
            self.SECCION,
            self.DATOS,
            [{"tipo": "foto_unica", "url": "https://x/z.webp",
              "checklist_item": "modulos.Campo Borrado"}],
        )
        assert generator.doc.inline_shapes


class TestFotosEnElDocumento:
    def test_la_foto_queda_embebida_con_su_pie(self, monkeypatch):
        monkeypatch.setattr(dg.httpx, "get", lambda url, **kw: _Respuesta(200, _png_bytes()))
        generator = DocxGenerator()
        generator.add_section(
            {"titulo": "Módulos", "icono": "", "campos": []},
            {},
            [{"tipo": "antes", "url": "https://x/f.webp", "checklist_item": "modulos.Limpieza"}],
        )

        assert generator.doc.inline_shapes  # la imagen entró al documento
        texto = "\n".join(p.text for p in generator.doc.paragraphs)
        texto += "\n".join(
            c.text for t in generator.doc.tables for f in t.rows for c in f.cells
        )
        assert "Ilustración 1. Vista de Limpieza antes" in texto

    def test_avisa_cuando_la_foto_no_se_puede_traer(self, monkeypatch):
        monkeypatch.setattr(dg.httpx, "get", lambda url, **kw: _Respuesta(403))
        generator = DocxGenerator()
        generator.add_section(
            {"titulo": "Módulos", "icono": "", "campos": []},
            {},
            [{"tipo": "antes", "url": "https://x/f.webp", "checklist_item": "modulos.Limpieza"}],
        )

        texto = "".join(
            c.text for t in generator.doc.tables for f in t.rows for c in f.cells
        )
        assert "Foto no disponible" in texto


class TestFormatearFecha:
    """El created_at del reporte llega como datetime desde SQLAlchemy"""

    def test_datetime(self):
        assert formatear_fecha(datetime(2026, 7, 24, 15, 30, 0)) == "2026-07-24"

    def test_date(self):
        assert formatear_fecha(date(2026, 7, 24)) == "2026-07-24"

    def test_texto_iso(self):
        assert formatear_fecha("2026-07-24T15:30:00Z") == "2026-07-24"

    @pytest.mark.parametrize("vacio", [None, ""])
    def test_sin_valor(self, vacio):
        assert formatear_fecha(vacio) == ""


class TestDocxGenerator:
    PROYECTO = {
        "nombre": "Planta Test",
        "cliente": "Cliente Test",
        "direccion": "Calle Falsa 123",
        "tipo_sistema": "on_grid",
        "potencia": "5.5",
    }

    def test_add_project_info_acepta_datetime(self):
        """Regresión: se hacía [:10] sobre el datetime en vez de sobre su texto,
        y el export fallaba con 500 (TypeError: not subscriptable)"""
        generator = DocxGenerator()
        reporte = {"created_at": datetime(2026, 7, 24), "tecnico_nombre": "Técnico Test"}

        generator.add_project_info(self.PROYECTO, reporte)

        texto = "\n".join(
            celda.text for tabla in generator.doc.tables for fila in tabla.rows for celda in fila.cells
        )
        assert "2026-07-24" in texto
        assert "Técnico Test" in texto

    def test_add_project_info_sin_fecha(self):
        generator = DocxGenerator()
        generator.add_project_info(self.PROYECTO, {"tecnico_nombre": ""})
        assert generator.generate()

    def test_export_completo_produce_un_docx(self):
        generator = DocxGenerator()
        generator.add_header_with_logo(self.PROYECTO)
        generator.add_project_info(
            self.PROYECTO, {"created_at": datetime(2026, 7, 24), "tecnico_nombre": "Técnico Test"}
        )
        generator.add_section(
            {"titulo": "Módulos", "icono": "☀️", "campos": [{"nombre": "Limpieza", "sin_fotos": True}]},
            {"Limpieza": "OK"},
            [],
        )

        contenido = generator.generate()
        # Un .docx es un zip: empieza con la firma PK
        assert contenido[:2] == b"PK"
        assert len(contenido) > 0
