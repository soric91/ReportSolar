from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from io import BytesIO
import base64
from datetime import date, datetime


def formatear_fecha(valor) -> str:
    """Fecha como YYYY-MM-DD, llegue como datetime o como texto ISO."""
    if not valor:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    return str(valor)[:10]


class DocxGenerator:
    def __init__(self, logo_url=None):
        self.logo_url = logo_url
        self.doc = Document()
        self.doc.sections[0].top_margin = Inches(0.5)
        self.doc.sections[0].bottom_margin = Inches(0.5)
        self.doc.sections[0].left_margin = Inches(0.75)
        self.doc.sections[0].right_margin = Inches(0.75)

    def add_header_with_logo(self, proyecto, logo_data=None):
        """Agrega encabezado con logo"""
        header_table = self.doc.add_table(rows=1, cols=2)
        header_table.autofit = False
        header_table.allow_autofit = False

        # Logo
        logo_cell = header_table.rows[0].cells[0]
        if logo_data:
            try:
                logo_cell.text = ""
                p = logo_cell.paragraphs[0]
                run = p.add_run()
                run.add_picture(BytesIO(logo_data), width=Inches(1.2))
            except:
                logo_cell.text = "PIVMAN SOLAR"

        # Título
        title_cell = header_table.rows[0].cells[1]
        title_cell.text = ""
        p = title_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("INFORME DE MANTENIMIENTO\n")
        run.bold = True
        run.font.size = Pt(16)
        run = p.add_run(f"Sistema Fotovoltaico\n")
        run.font.size = Pt(12)
        run = p.add_run(f"{proyecto.get('nombre', 'N/A')}")
        run.font.size = Pt(11)

        # Línea separadora
        self._add_separator()

    def add_project_info(self, proyecto, reporte):
        """Agrega datos del proyecto"""
        p = self.doc.add_paragraph()
        p.style = "Heading 2"
        run = p.add_run("Datos del Proyecto")
        run.font.size = Pt(12)

        table = self.doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"

        data = [
            ("Cliente", proyecto.get("cliente", "")),
            ("Dirección", proyecto.get("direccion", "")),
            ("Tipo de Sistema", proyecto.get("tipo_sistema", "")),
            ("Potencia (kW)", proyecto.get("potencia", "")),
            ("Fecha de Visita", formatear_fecha(reporte.get("created_at"))),
            ("Técnico", reporte.get("tecnico_nombre", "")),
        ]

        for label, value in data:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[1].text = str(value)

        self.doc.add_paragraph()

    def add_section(self, section_def, section_data, fotos):
        """Agrega una sección con sus campos y fotos"""
        if not section_data or len(section_data) == 0:
            return

        # Título de sección
        p = self.doc.add_paragraph()
        p.style = "Heading 2"
        run = p.add_run(f"{section_def.get('icono', '')} {section_def.get('titulo', '')}")
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(200, 100, 0)  # Pivman orange

        # Campos
        for campo in section_def.get("campos", []):
            if campo.get("sin_fotos"):
                valor = section_data.get(campo.get("nombre", ""), "")
                if valor:
                    self._add_field(campo.get("nombre", ""), valor)

        # Fotos
        self._add_photos_for_section(section_data, fotos, section_def)

        self.doc.add_paragraph()

    def _add_field(self, label, value):
        """Agrega un campo con valor"""
        p = self.doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        p.add_run(str(value))

    def _add_photos_for_section(self, section_data, fotos, section_def):
        """Agrega fotos de una sección"""
        if not fotos or len(fotos) == 0:
            return

        foto_table = self.doc.add_table(rows=1, cols=3)
        foto_table.style = "Light Grid Accent 1"

        # Headers
        for cell, header in zip(foto_table.rows[0].cells, ["Antes", "Durante", "Después"]):
            cell.text = header

        # Fotos por tipo
        tipos = ["antes", "durante", "despues"]
        row = foto_table.add_row()

        for col_idx, tipo in enumerate(tipos):
            cell = row.cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]

            fotos_tipo = [f for f in fotos if f.get("tipo") == tipo]
            for foto in fotos_tipo[:1]:  # Una foto por tipo
                try:
                    if foto.get("url", "").startswith("data:"):
                        img_data = base64.b64decode(foto["url"].split(",")[1])
                    else:
                        img_data = foto["url"]
                    run = p.add_run()
                    run.add_picture(BytesIO(img_data), width=Inches(1.5))
                    p.add_run("\n")
                except:
                    p.add_run(f"No se pudo cargar foto ({tipo})")

    def _add_separator(self):
        """Agrega línea separadora"""
        p = self.doc.add_paragraph()
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "c86400")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def generate(self):
        """Retorna bytes del documento DOCX"""
        output = BytesIO()
        self.doc.save(output)
        output.seek(0)
        return output.getvalue()
